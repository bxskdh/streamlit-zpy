import streamlit as st
import numpy as np
from PIL import Image, ImageDraw
import os
import io
from datetime import datetime
from streamlit_paste_button import paste_image_button as pbutton

import sys
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from utils.gel_image import (
    detect_lanes_smart, detect_bands_in_lane, detect_bands_from_reference,
    map_bands_to_kda, create_annotated_image, pil_to_cv2, cv2_to_pil,
    get_projection_viz, draw_lane_overlay, get_band_projection_viz,
    ocr_marker_numbers
)
from utils.gel_db import (
    load_markers, save_custom_marker, save_preset_marker, save_annotation,
    get_next_default_name, GEL_IMAGES_DIR
)

import io
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from utils.gel_image import (
    detect_lanes_smart, detect_bands_in_lane, detect_bands_from_reference,
    map_bands_to_kda, create_annotated_image, pil_to_cv2, cv2_to_pil,
    get_projection_viz, draw_lane_overlay, get_band_projection_viz,
    ocr_marker_numbers, get_font
)

def create_annotated_pptx(
    original_image,
    lane_names,
    lane_boxes,
    marker_lane_idx,
    marker_bands_data,
    font_name="Arial",
    lane_label_y_offset=100,
    lane_label_x_offset=0,
    lane_label_angle=0,
    kda_label_x_offset=10,
    kda_label_y_offset=0,
    lane_fine_offsets=None,
    band_fine_offsets=None,
):
    """生成可编辑的 PPTX：图片作为底层，所有标注为独立文本框"""
    lane_fine_offsets = lane_fine_offsets or {}
    band_fine_offsets = band_fine_offsets or {}

    orig_w, orig_h = original_image.size
    top_pad = 120
    right_pad = 250
    left_pad = 80
    new_w = orig_w + left_pad + right_pad
    new_h = orig_h + top_pad

    EMU_PER_PX = 12700  # 72 dpi

    prs = Presentation()
    prs.slide_width = Emu(int(new_w * EMU_PER_PX))
    prs.slide_height = Emu(int(new_h * EMU_PER_PX))

    blank_layout = None
    for layout in prs.slide_layouts:
        if layout.name == 'Blank':
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = prs.slide_layouts[-1]
    slide = prs.slides.add_slide(blank_layout)

    # 底层：原图
    img_buf = io.BytesIO()
    original_image.save(img_buf, format='PNG')
    img_buf.seek(0)
    slide.shapes.add_picture(
        img_buf,
        Emu(int(left_pad * EMU_PER_PX)),
        Emu(int(top_pad * EMU_PER_PX)),
        width=Emu(int(orig_w * EMU_PER_PX)),
        height=Emu(int(orig_h * EMU_PER_PX))
    )

    # 胶道名称文本框
    lane_font_size = max(16, min(28, orig_w // 35))
    for i, (box, name) in enumerate(zip(lane_boxes, lane_names)):
        if not name:
            continue
        x1_orig, x2_orig = box
        x1 = x1_orig + left_pad
        x2 = x2_orig + left_pad

        fine = lane_fine_offsets.get(i, {})
        fine_x = fine.get("x", 0)
        fine_y = fine.get("y", 0)

        center_x = (x1 + x2) / 2 + fine_x + lane_label_x_offset
        center_y = lane_label_y_offset + fine_y

        pil_font = get_font(font_name, lane_font_size)
        draw_tmp = ImageDraw.Draw(Image.new('RGB', (1, 1)))
        bbox = draw_tmp.textbbox((0, 0), name, font=pil_font)
        text_w = bbox[2] - bbox[0]
        text_h = bbox[3] - bbox[1]

        tx_left = center_x - text_w / 2
        tx_top = center_y - text_h / 2

        txBox = slide.shapes.add_textbox(
            Emu(int(tx_left * EMU_PER_PX)),
            Emu(int(tx_top * EMU_PER_PX)),
            Emu(int(text_w * EMU_PER_PX)),
            Emu(int(text_h * EMU_PER_PX))
        )
        tf = txBox.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = name
        p.font.size = Pt(lane_font_size)
        p.font.name = font_name.split()[0]
        p.font.color.rgb = RGBColor(0, 0, 0)
        p.alignment = PP_ALIGN.CENTER

        if lane_label_angle != 0:
            txBox.rotation = lane_label_angle

    # Marker kDa 文本框
    if marker_lane_idx is not None and marker_bands_data and 0 <= marker_lane_idx < len(lane_boxes):
        lane_box = lane_boxes[marker_lane_idx]
        x1_orig, x2_orig = lane_box
        x1 = x1_orig + left_pad

        kda_font_size = max(12, min(20, orig_w // 45))

        for bidx, band in enumerate(marker_bands_data):
            y_center_orig = band["y"]
            y_center = y_center_orig + top_pad
            kda = band["kda"]
            unit = band.get("unit", "kDa")
            label = f"{kda}{unit}" if unit != "kDa" else f"{kda}"

            fine = band_fine_offsets.get(bidx, {})
            fine_x = fine.get("x", 0)
            fine_y = fine.get("y", 0)

            pil_font = get_font(font_name, kda_font_size)
            draw_tmp = ImageDraw.Draw(Image.new('RGB', (1, 1)))
            bbox = draw_tmp.textbbox((0, 0), label, font=pil_font)
            text_w = bbox[2] - bbox[0]
            text_h = bbox[3] - bbox[1]

            label_x = x1 - kda_label_x_offset - text_w + fine_x
            label_y = y_center - text_h / 2 + fine_y + kda_label_y_offset

            txBox = slide.shapes.add_textbox(
                Emu(int(label_x * EMU_PER_PX)),
                Emu(int(label_y * EMU_PER_PX)),
                Emu(int(text_w * EMU_PER_PX)),
                Emu(int(text_h * EMU_PER_PX))
            )
            tf = txBox.text_frame
            tf.word_wrap = False
            p = tf.paragraphs[0]
            p.text = label
            p.font.size = Pt(kda_font_size)
            p.font.name = font_name.split()[0]
            p.font.color.rgb = RGBColor(0, 0, 0)
            p.alignment = PP_ALIGN.LEFT

    out_buf = io.BytesIO()
    prs.save(out_buf)
    out_buf.seek(0)
    return out_buf

st.set_page_config(page_title="胶图自动标注", layout="wide")

st.title("🧬 凝胶电泳图自动标注系统")

# ========== Session State ==========
defaults = {
    "gel_image": None,
    "lane_boxes": [],
    "lane_names": [],
    "marker_lane_idx": None,
    "detected_bands": [],
    "band_kda_mapping": {},
    "annotation_name": "",
    "mapped_bands": [],
    "expected_lanes": None,
    "auto_expected": None,
    "projection": None,
    "peaks": None,
    "valleys": None,
    "last_annotated": None,
    "manual_offset": 0,
    "selected_marker_source": "preset",
    "upload_marker_bands": [],
    "ref_band_kdas": [],
    "expected_bands": None,
    "band_sensitivity": 0.25,
    "band_detection_result": None,
    # === 双模式胶道识别 ===
    "method_a_lanes": [],
    "method_a_projection": None,
    "method_a_peaks": None,
    "method_a_valleys": None,
    "method_a_offset": 0,
    "method_a_expected": None,
    "method_b_lanes": [],
    "method_b_projection": None,
    "method_b_valleys": None,
    "method_b_all_peaks": [],
    "method_b_offset": 0,
    "method_b_expected": None,
    "selected_lane_method": "A",
    "global_x_shift": 0,
    "lane_individual_shifts": {},
}
for k, v in defaults.items():
    if k not in st.session_state:
        st.session_state[k] = v

if "lane_fine_offsets" not in st.session_state:
    st.session_state.lane_fine_offsets = {}
if "band_fine_offsets" not in st.session_state:
    st.session_state.band_fine_offsets = {}


def auto_detect_marker_idx(names):
    for i, name in enumerate(names):
        lower = name.lower()
        if "mk" in lower or "marker" in lower:
            return i
    return None

def create_lanes_from_peaks(peaks, valleys, img_w, min_width=15):
    """根据峰和谷创建胶道框，用于峰值识别模式"""
    if not peaks:
        return []
    peaks = sorted(peaks)
    valleys = sorted(valleys)
    lanes = []
    for i, p in enumerate(peaks):
        left_candidates = [v for v in valleys if v < p]
        if left_candidates:
            left = left_candidates[-1]
        else:
            left = max(0, p - min_width)

        right_candidates = [v for v in valleys if v > p]
        if right_candidates:
            right = right_candidates[0]
        else:
            right = min(img_w, p + min_width)

        if right - left < min_width:
            center = (left + right) / 2
            left = int(center - min_width / 2)
            right = int(center + min_width / 2)

        lanes.append((int(left), int(right)))
    return lanes

# ========== 侧边栏 ==========
with st.sidebar:
    st.header("⚙️ 全局参数")
    
    font_options = ["Arial", "Times New Roman", "Courier New",
                   "SimHei (黑体)", "SimSun (宋体)", "Microsoft YaHei (微软雅黑)"]
    selected_font = st.selectbox("字体", font_options, index=0)
    
    st.divider()
    st.subheader("胶道标注统一参数")
    lane_label_y_offset = st.slider("统一Y轴偏移 (紧贴原图上方为100)", 40, 150, 100)
    lane_label_x_offset = st.slider("统一X轴偏移 (像素，可负)", -100, 100, 0)
    lane_label_angle = st.slider("统一倾斜角度", -90, 90, 0)
    
    st.divider()
    st.subheader("MK标注统一参数")
    kda_label_x_offset = st.slider("统一X轴偏移 (像素，可负)", -100, 100, 10)
    kda_label_y_offset = st.slider("统一Y轴偏移 (像素，可负)", -100, 100, 0)
    
    show_preview = st.checkbox("实时预览", value=True)
    show_lane_boxes = st.checkbox("显示胶道框", value=False)

# ========== 主界面 ==========
tab1, tab2, tab3, tab4 = st.tabs(["📤 上传与胶道识别", "🔬 Marker校准", "💾 预览与保存", "📝 标注管理"])

# ========== Tab 1: 上传与识别 ==========
with tab1:
    st.subheader("1. 上传凝胶电泳图像")
    uploaded_file = st.file_uploader("支持 JPG, PNG, TIFF 格式",
                                    type=["jpg", "jpeg", "png", "tif", "tiff"])
    
    # ---- 从剪贴板粘贴图片 ----
    st.divider()
    col_paste, _ = st.columns([1, 3])
    with col_paste:
        paste_result = pbutton(
            label="📋 从剪贴板粘贴图片（Ctrl+V）",
            key="paste_button_auto",
            errors="raise"
        )
        if paste_result.image_data is not None:
            pil_img = paste_result.image_data
            st.session_state.gel_image = pil_img
            st.session_state.crop_left = 0
            st.session_state.crop_right = 0
            st.session_state.crop_top = 0
            st.session_state.crop_bottom = 0
            st.session_state.lane_boxes = []
            st.session_state.lane_names = []
            st.session_state.marker_lane_idx = None
            st.session_state.detected_bands = []
            st.session_state.mapped_bands = []
            st.session_state.method_a_lanes = []
            st.session_state.method_b_lanes = []
            st.session_state.band_detection_result = None
            st.success("✅ 图片粘贴成功！")
            st.rerun()
    
    # 文件上传处理
    if uploaded_file is not None:
        image = Image.open(uploaded_file).convert("RGB")
        st.session_state.gel_image = image
        st.session_state.crop_left = 0
        st.session_state.crop_right = 0
        st.session_state.crop_top = 0
        st.session_state.crop_bottom = 0
        st.session_state.lane_boxes = []
        st.session_state.lane_names = []
        st.session_state.marker_lane_idx = None
        st.session_state.detected_bands = []
        st.session_state.mapped_bands = []
        st.session_state.method_a_lanes = []
        st.session_state.method_b_lanes = []
        st.session_state.band_detection_result = None
        st.rerun()
    
    # ---- 如果已有图片，显示后续所有界面 ----
    if st.session_state.gel_image is not None:
        image = st.session_state.gel_image
        cv2_img = pil_to_cv2(image)   # 确保已导入
        img_w, img_h = image.size
        
        # ========== 图片裁剪 ==========
        st.divider()
        st.subheader("1.5 图片裁剪（去除四周空白）")
        
        # 初始化裁剪值（确保在 session_state 中）
        if 'crop_left' not in st.session_state:
            st.session_state.crop_left = 0
            st.session_state.crop_right = 0
            st.session_state.crop_top = 0
            st.session_state.crop_bottom = 0
        
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            st.session_state.crop_left = st.number_input("左侧裁剪 (px)", 0, 99999, st.session_state.crop_left)
        with c2:
            st.session_state.crop_right = st.number_input("右侧裁剪 (px)", 0, 99999, st.session_state.crop_right)
        with c3:
            st.session_state.crop_top = st.number_input("上侧裁剪 (px)", 0, 99999, st.session_state.crop_top)
        with c4:
            st.session_state.crop_bottom = st.number_input("下侧裁剪 (px)", 0, 99999, st.session_state.crop_bottom)
        
        if (st.session_state.crop_left + st.session_state.crop_right < img_w and
            st.session_state.crop_top + st.session_state.crop_bottom < img_h):
            cropped = image.crop((
                st.session_state.crop_left,
                st.session_state.crop_top,
                img_w - st.session_state.crop_right,
                img_h - st.session_state.crop_bottom
            ))
            st.image(cropped, caption="裁剪预览", width='stretch')
            st.session_state.gel_image = cropped
            cv2_img = pil_to_cv2(cropped)
            image = cropped
            st.success(f"已裁剪: 原图 {img_w}x{img_h} → {cropped.width}x{cropped.height}")
        else:
            st.error("裁剪值过大，请调整")
        
        col_exp1, col_exp2 = st.columns([1, 2])
        with col_exp1:
            use_auto = st.checkbox("自动推断胶道数", value=(st.session_state.expected_lanes is None))
        with col_exp2:
            if use_auto:
                st.session_state.expected_lanes = None
                st.markdown("🤖 **自动模式**")
            else:
                user_expected = st.number_input(
                    "预期胶道数", 1, 30,
                    value=st.session_state.expected_lanes if st.session_state.expected_lanes else 9
                )
                st.session_state.expected_lanes = int(user_expected)
                st.markdown(f"📍 **固定模式**：强制分割为 **{st.session_state.expected_lanes}** 条等宽胶道")
        
        # ========== 胶道识别方法 ==========
        st.divider()
        st.subheader("3. 胶道识别方法")
        st.info("""
        提供两种胶道识别方法，请分别调整参数后查看结果，最后选择一种进入后续步骤。
        - **等距分割**：基于水平投影均匀分割胶道
        - **峰值识别**：仅将检测到的峰（有内容的区域）作为胶道，可手动增删
        """)
        
        method_a_col, method_b_col = st.columns(2)
        
        # ========== 方法A：等距分割 ==========
        with method_a_col:
            st.markdown("**方法A：等距分割**")
            use_auto_a = st.checkbox("自动推断胶道数", value=(st.session_state.method_a_expected is None), key="auto_a")
            if use_auto_a:
                st.session_state.method_a_expected = None
                st.markdown("🤖 **自动模式**")
            else:
                val_a = st.session_state.method_a_expected if st.session_state.method_a_expected else 9
                user_expected_a = st.number_input("预期胶道数", 1, 30, value=val_a, key="exp_a")
                st.session_state.method_a_expected = int(user_expected_a)
                st.markdown(f"📍 **固定模式**：强制分割为 **{st.session_state.method_a_expected}** 条等宽胶道")
            lane_width_a = img_w / max(st.session_state.method_a_expected if st.session_state.method_a_expected else 9, 1)
            max_offset_a = int(lane_width_a / 2)
            st.session_state.method_a_offset = st.slider(
                "水平偏移（像素，正值=右移）", -max_offset_a, max_offset_a,
                value=st.session_state.method_a_offset, step=1, key="offset_a"
            )
            if st.button("🔄 重置偏移", key="reset_a"):
                st.session_state.method_a_offset = 0
                st.rerun()
            if st.button("🔍 检测胶道（等距分割）", type="primary", key="detect_a", use_container_width=True):
                lanes_a, proj_a, peaks_a, valleys_a = detect_lanes_smart(
                    cv2_img,
                    expected_lanes=st.session_state.method_a_expected,
                    min_lane_width=15,
                    manual_offset=st.session_state.method_a_offset
                )
                st.session_state.method_a_lanes = lanes_a
                st.session_state.method_a_projection = proj_a
                st.session_state.method_a_peaks = peaks_a
                st.session_state.method_a_valleys = valleys_a
                st.success(f"检测到 {len(lanes_a)} 条胶道")
                st.rerun()
            if st.session_state.method_a_lanes:
                proj_viz_a = get_projection_viz(
                    cv2_img, st.session_state.method_a_lanes,
                    st.session_state.method_a_projection,
                    st.session_state.method_a_peaks,
                    st.session_state.method_a_valleys,
                    expected_lanes=st.session_state.method_a_expected
                )
                st.image(cv2_to_pil(proj_viz_a), caption="📊 水平投影分析图", width='stretch')
                overlay_a = draw_lane_overlay(
                    image, st.session_state.method_a_lanes,
                    selected_idx=None,
                    peaks=st.session_state.method_a_peaks,
                    valleys=st.session_state.method_a_valleys
                )
                st.image(overlay_a, caption=f"🎨 胶道分布图 (共{len(st.session_state.method_a_lanes)}条)", width='stretch')
                st.metric("检测到的胶道数", len(st.session_state.method_a_lanes))
                st.metric("检测到的峰数", len(st.session_state.method_a_peaks))
                if len(st.session_state.method_a_peaks) < len(st.session_state.method_a_lanes):
                    st.info(f"ℹ️ {len(st.session_state.method_a_peaks)}个有内容的峰 + {len(st.session_state.method_a_lanes)-len(st.session_state.method_a_peaks)}个空白胶道")
        
        # ========== 方法B：峰值识别 ==========
        with method_b_col:
            st.markdown("**方法B：峰值识别**")
            val_b = st.session_state.method_b_expected if st.session_state.method_b_expected else 9
            st.session_state.method_b_expected = st.number_input("预期胶道数", 1, 30, value=val_b, key="exp_b")
            max_offset_b = int(img_w / 2 / max(st.session_state.method_b_expected, 1))
            st.session_state.method_b_offset = st.slider(
                "水平偏移（像素，正值=右移）", -max_offset_b, max_offset_b,
                value=st.session_state.method_b_offset, step=1, key="offset_b"
            )
            if st.button("🔄 重置偏移", key="reset_b"):
                st.session_state.method_b_offset = 0
                st.rerun()
            if st.button("🔍 检测胶道（峰值识别）", type="primary", key="detect_b", use_container_width=True):
                _, proj_b, peaks_b, valleys_b = detect_lanes_smart(
                    cv2_img,
                    expected_lanes=None,
                    min_lane_width=15,
                    manual_offset=st.session_state.method_b_offset
                )
                st.session_state.method_b_projection = proj_b
                st.session_state.method_b_valleys = valleys_b
                st.session_state.method_b_all_peaks = sorted(list(peaks_b))
                st.session_state.method_b_lanes = create_lanes_from_peaks(
                    st.session_state.method_b_all_peaks, valleys_b, img_w
                )
                st.success(f"检测到 {len(peaks_b)} 个峰，生成 {len(st.session_state.method_b_lanes)} 条胶道")
                st.rerun()
            if st.session_state.method_b_lanes:
                proj_viz_b = get_projection_viz(
                    cv2_img, st.session_state.method_b_lanes,
                    st.session_state.method_b_projection,
                    st.session_state.method_b_all_peaks,
                    st.session_state.method_b_valleys,
                    expected_lanes=None
                )
                st.image(cv2_to_pil(proj_viz_b), caption="📊 水平投影分析图", width='stretch')
                overlay_b = draw_lane_overlay(
                    image, st.session_state.method_b_lanes,
                    selected_idx=None,
                    peaks=st.session_state.method_b_all_peaks,
                    valleys=st.session_state.method_b_valleys
                )
                st.image(overlay_b, caption=f"🎨 胶道分布图 (共{len(st.session_state.method_b_lanes)}条)", width='stretch')
                n_peaks = len(st.session_state.method_b_all_peaks)
                expected = st.session_state.method_b_expected
                if n_peaks < expected:
                    st.warning(f"⚠️ 当前 {n_peaks} 条胶道，少于预期的 {expected} 条，需拆分现有胶道")
                    st.info("点击下方胶道的「拆分」按钮，输入分割点x坐标（可多个，用逗号分隔）")
                    n_show = min(n_peaks, 6)
                    split_cols = st.columns(n_show)
                    selected_split_idx = st.session_state.get("method_b_split_idx", None)
                    for idx in range(n_peaks):
                        col = split_cols[idx % n_show]
                        with col:
                            x1, x2 = st.session_state.method_b_lanes[idx]
                            lane_crop = image.crop((x1, 0, x2, image.height))
                            st.image(lane_crop, width=20)
                            if st.button(f"拆分 #{idx+1}", key=f"split_{idx}", use_container_width=True):
                                st.session_state.method_b_split_idx = idx
                                st.rerun()
                    if selected_split_idx is not None:
                        st.divider()
                        st.markdown(f"**拆分胶道 {selected_split_idx+1}**")
                        x1, x2 = st.session_state.method_b_lanes[selected_split_idx]
                        st.image(image.crop((x1, 0, x2, image.height)), caption=f"当前范围: x={x1}~{x2}", width=62)
                        split_points_str = st.text_input(
                            f"分割点x坐标（{x1}~{x2}，多个用逗号分隔）",
                            placeholder=f"例如: {x1 + (x2-x1)//3}, {x1 + 2*(x2-x1)//3}",
                            key="split_points"
                        )
                        c1, c2 = st.columns(2)
                        with c1:
                            if st.button("✅ 确认拆分", key="confirm_split", use_container_width=True):
                                try:
                                    points = [int(float(p.strip())) for p in split_points_str.split(",") if p.strip()]
                                    points = sorted(list(set([p for p in points if x1 < p < x2])))
                                    if not points:
                                        st.error("请输入有效的分割点")
                                    else:
                                        old_lane = st.session_state.method_b_lanes[selected_split_idx]
                                        boundaries = [old_lane[0]] + points + [old_lane[1]]
                                        new_lanes = []
                                        new_peaks = []
                                        for i in range(len(boundaries) - 1):
                                            new_lanes.append((boundaries[i], boundaries[i+1]))
                                            new_peaks.append((boundaries[i] + boundaries[i+1]) // 2)
                                        st.session_state.method_b_lanes = (
                                            st.session_state.method_b_lanes[:selected_split_idx] + 
                                            new_lanes + 
                                            st.session_state.method_b_lanes[selected_split_idx+1:]
                                        )
                                        st.session_state.method_b_all_peaks = (
                                            st.session_state.method_b_all_peaks[:selected_split_idx] + 
                                            new_peaks + 
                                            st.session_state.method_b_all_peaks[selected_split_idx+1:]
                                        )
                                        st.session_state.method_b_split_idx = None
                                        st.success(f"已拆分为 {len(new_lanes)} 条胶道")
                                        st.rerun()
                                except Exception as e:
                                    st.error(f"拆分失败: {e}")
                        with c2:
                            if st.button("❌ 取消", key="cancel_split", use_container_width=True):
                                st.session_state.method_b_split_idx = None
                                st.rerun()
                elif n_peaks > expected:
                    st.warning(f"⚠️ 当前 {n_peaks} 条胶道，多于预期的 {expected} 条，请删除多余胶道")
                    del_options = [f"胶道 {i+1} (x≈{int(p)})" for i, p in enumerate(st.session_state.method_b_all_peaks)]
                    del_idx = st.selectbox("选择要删除的胶道", range(len(del_options)), 
                                          format_func=lambda x: del_options[x], key="del_b")
                    if st.button("🗑️ 删除选中胶道", key="del_lane_b", use_container_width=True):
                        st.session_state.method_b_all_peaks.pop(del_idx)
                        st.session_state.method_b_lanes = create_lanes_from_peaks(
                            st.session_state.method_b_all_peaks, st.session_state.method_b_valleys, img_w
                        )
                        st.rerun()
                else:
                    st.success(f"✅ 当前 {n_peaks} 条胶道，符合预期")
        
        # ========== 选择识别结果 ==========
        st.divider()
        st.subheader("4. 选择识别结果")
        has_a = len(st.session_state.method_a_lanes) > 0
        has_b = len(st.session_state.method_b_lanes) > 0
        options = []
        if has_a:
            options.append(f"方法A（等距分割）- {len(st.session_state.method_a_lanes)}条胶道")
        if has_b:
            options.append(f"方法B（峰值识别）- {len(st.session_state.method_b_lanes)}条胶道")
        if not options:
            st.error("⚠️ 请先在上方至少一种方法中检测胶道")
            st.stop()
        default_idx = 0 if st.session_state.selected_lane_method == "A" or not has_b else 1
        selected = st.radio("选择要使用的胶道识别结果", options, index=default_idx)
        st.session_state.selected_lane_method = "A" if "方法A" in selected else "B"
        
        # ========== 全局平移 ==========
        st.divider()
        st.subheader("5. 全局X轴平移")
        st.info("对所有已选胶道进行统一的左右平移，以精确对准胶道位置")
        
        if st.session_state.selected_lane_method == "A":
            base_lanes = st.session_state.method_a_lanes
            base_peaks = st.session_state.method_a_peaks if st.session_state.method_a_peaks is not None else []
            base_valleys = st.session_state.method_a_valleys if st.session_state.method_a_valleys is not None else []
        else:
            base_lanes = st.session_state.method_b_lanes
            base_peaks = list(st.session_state.method_b_all_peaks) if st.session_state.method_b_all_peaks else []
            base_valleys = st.session_state.method_b_valleys if st.session_state.method_b_valleys is not None else []
        
        max_shift_left = -img_w
        max_shift_right = img_w
        if base_lanes:
            first_lane = base_lanes[0]
            last_lane = base_lanes[-1]
            margin_left = first_lane[0]
            margin_right = img_w - last_lane[1]
            st.caption(f"胶道边界余量：左 {margin_left}px，右 {margin_right}px（超出范围请自行判断）")
        
        st.session_state.global_x_shift = st.slider(
            "全局X轴平移（像素，正值=右移，负值=左移）",
            max_shift_left, max_shift_right,
            value=st.session_state.global_x_shift,
            step=1
        )
        if st.button("🔄 重置平移为0", key="reset_global_shift"):
            st.session_state.global_x_shift = 0
            st.rerun()
        
        # 应用平移
        shift = st.session_state.global_x_shift
        st.session_state.lane_boxes = [(x1+shift, x2+shift) for x1, x2 in base_lanes]
        st.session_state.projection = st.session_state.method_a_projection if st.session_state.selected_lane_method == "A" else st.session_state.method_b_projection
        st.session_state.peaks = [p+shift for p in base_peaks]
        st.session_state.valleys = [v+shift for v in base_valleys]
        
        st.markdown("**平移后预览：**")
        viz_col1, viz_col2 = st.columns([2, 1])
        with viz_col1:
            proj_viz = get_projection_viz(
                cv2_img, st.session_state.lane_boxes,
                st.session_state.projection,
                st.session_state.peaks,
                st.session_state.valleys,
                expected_lanes=st.session_state.method_a_expected if st.session_state.selected_lane_method == "A" else None
            )
            st.image(cv2_to_pil(proj_viz), caption="📊 水平投影分析图（平移后）| 绿色=峰 | 红色=谷 | 橙框=胶道", width='stretch')
        with viz_col2:
            overlay_img = draw_lane_overlay(
                image, st.session_state.lane_boxes,
                selected_idx=st.session_state.marker_lane_idx,
                peaks=st.session_state.peaks, valleys=st.session_state.valleys
            )
            st.image(overlay_img, caption=f"🎨 胶道分布图 (共{len(st.session_state.lane_boxes)}条)", width='stretch')
            st.metric("当前胶道数", len(st.session_state.lane_boxes))
        
        # ========== 胶道命名、微调与实时预览 ==========
        st.divider()
        st.subheader("5. 胶道命名、微调与实时预览")
        
        # 计算 final_lanes（global_shift + individual_shift）
        shift = st.session_state.global_x_shift
        base_with_global = [(x1 + shift, x2 + shift) for x1, x2 in base_lanes]
        final_lanes = []
        for i, (x1, x2) in enumerate(base_with_global):
            ind_shift = st.session_state.lane_individual_shifts.get(i, 0)
            final_lanes.append((x1 + ind_shift, x2 + ind_shift))
        st.session_state.lane_boxes = final_lanes
        st.session_state.peaks = [p + shift for p in base_peaks]
        st.session_state.valleys = [v + shift for v in base_valleys]
        
        if len(st.session_state.lane_names) != len(st.session_state.lane_boxes):
            old_names = st.session_state.lane_names
            new_names = []
            for i in range(len(st.session_state.lane_boxes)):
                if i < len(old_names):
                    new_names.append(old_names[i])
                else:
                    new_names.append(f"Lane {i+1}")
            st.session_state.lane_names = new_names
        
        auto_mk = auto_detect_marker_idx(st.session_state.lane_names)
        if auto_mk is not None and st.session_state.marker_lane_idx is None:
            st.session_state.marker_lane_idx = auto_mk
        
        left_col, right_col = st.columns([1, 2])
        with left_col:
            st.info("🗑️ 名称含MK会自动设为Marker。每个胶道可独立X轴偏移。点击删除移除胶道。")
            n_lanes = len(st.session_state.lane_boxes)
            lanes_per_row = 3
            for i in range(len(st.session_state.lane_names)):
                key = f"lane_name_{i}"
                if key in st.session_state:
                    st.session_state[key] = st.session_state.lane_names[i]
            for row_start in range(0, n_lanes, lanes_per_row):
                cols = st.columns(lanes_per_row)
                for j, col in enumerate(cols):
                    i = row_start + j
                    if i >= n_lanes:
                        break
                    with col:
                        x1, x2 = st.session_state.lane_boxes[i]
                        x1 = max(0, x1)
                        x2 = min(image.width, x2)
                        if x2 > x1:
                            lane_crop = image.crop((x1, 0, x2, image.height))
                            st.image(lane_crop, width=20)
                        else:
                            st.error("范围错误")
                        default_name = st.session_state.lane_names[i] if i < len(st.session_state.lane_names) else f"Lane {i+1}"
                        name = st.text_input(f"名称", value=default_name, key=f"lane_name_{i}", label_visibility="collapsed")
                        if i < len(st.session_state.lane_names):
                            st.session_state.lane_names[i] = name
                        current_ind_shift = st.session_state.lane_individual_shifts.get(i, 0)
                        ind_shift = st.number_input("X偏移", value=current_ind_shift, step=1, key=f"ind_shift_{i}")
                        st.session_state.lane_individual_shifts[i] = ind_shift
                        if st.button("🗑️", key=f"del_lane_{i}", help="删除此胶道"):
                            if st.session_state.selected_lane_method == "A":
                                if i < len(st.session_state.method_a_lanes):
                                    st.session_state.method_a_lanes.pop(i)
                            else:
                                if i < len(st.session_state.method_b_lanes):
                                    st.session_state.method_b_lanes.pop(i)
                                if i < len(st.session_state.method_b_all_peaks):
                                    st.session_state.method_b_all_peaks.pop(i)
                            if i in st.session_state.lane_individual_shifts:
                                del st.session_state.lane_individual_shifts[i]
                            new_shifts = {}
                            for k, v in st.session_state.lane_individual_shifts.items():
                                if k > i:
                                    new_shifts[k - 1] = v
                                elif k < i:
                                    new_shifts[k] = v
                            st.session_state.lane_individual_shifts = new_shifts
                            if i < len(st.session_state.lane_names):
                                st.session_state.lane_names.pop(i)
                            if st.session_state.marker_lane_idx == i:
                                remaining = len(st.session_state.method_a_lanes if st.session_state.selected_lane_method == "A" else st.session_state.method_b_lanes)
                                st.session_state.marker_lane_idx = 0 if remaining > 0 else None
                            elif st.session_state.marker_lane_idx is not None and st.session_state.marker_lane_idx > i:
                                st.session_state.marker_lane_idx -= 1
                            st.rerun()
        
        with right_col:
            st.markdown("**最终胶道分布预览（实时）**")
            proj_viz = get_projection_viz(
                cv2_img, st.session_state.lane_boxes,
                st.session_state.projection,
                st.session_state.peaks,
                st.session_state.valleys,
                expected_lanes=st.session_state.method_a_expected if st.session_state.selected_lane_method == "A" else None
            )
            st.image(cv2_to_pil(proj_viz), caption="📊 水平投影分析图（含所有偏移）", width='stretch')
            overlay_img = draw_lane_overlay(
                image, st.session_state.lane_boxes,
                selected_idx=st.session_state.marker_lane_idx,
                peaks=st.session_state.peaks, valleys=st.session_state.valleys
            )
            st.image(overlay_img, caption=f"🎨 胶道分布图 (共{len(st.session_state.lane_boxes)}条)", width='stretch')
        
        # ========== Marker胶道选择 ==========
        st.divider()
        st.subheader("6. Marker胶道选择")
        lane_options = [f"胶道 {i+1}: {st.session_state.lane_names[i][:15]}"
                      for i in range(len(st.session_state.lane_boxes))]
        if lane_options:
            marker_idx = st.selectbox(
                "选择 Marker 胶道",
                options=list(range(len(lane_options))),
                format_func=lambda x: lane_options[x],
                index=st.session_state.marker_lane_idx if st.session_state.marker_lane_idx is not None else 0
            )
            st.session_state.marker_lane_idx = marker_idx
        
        # ========== Marker条带检测 ==========
        st.divider()
        st.subheader("7. Marker条带检测")
        st.info("""
        💡 **检测原理**：计算Marker胶道的**垂直投影**（每行亮度），反转后峰=条带。
        绿色圆点=检测到的峰（条带中心），红色圆点=谷（条带间边界），橙框=最终条带区域。
        """)
        
        band_col1, band_col2, band_col3 = st.columns(3)
        with band_col1:
            st.session_state.band_sensitivity = st.slider(
                "检测灵敏度", 0.05, 0.8, st.session_state.band_sensitivity, 0.05,
                help="越低越敏感，能检测更弱的条带"
            )
        with band_col2:
            st.session_state.expected_bands = st.number_input(
                "预期条带数（可选）", 0, 30,
                value=st.session_state.expected_bands if st.session_state.expected_bands else 0,
                help="设为0则自动检测；设为具体数字则强制按此数量均匀分割条带"
            )
            if st.session_state.expected_bands == 0:
                st.session_state.expected_bands = None
        with band_col3:
            st.markdown("&nbsp;")
            detect_btn = st.button("🔍 检测Marker条带", type="primary", use_container_width=True)
        
        if detect_btn:
            if st.session_state.marker_lane_idx is not None and st.session_state.marker_lane_idx < len(st.session_state.lane_boxes):
                lane_box = st.session_state.lane_boxes[st.session_state.marker_lane_idx]
                exp_bands = st.session_state.expected_bands if st.session_state.expected_bands and st.session_state.expected_bands > 0 else None
                result = detect_bands_in_lane(
                    cv2_img, lane_box,
                    min_band_height=2,
                    sensitivity=st.session_state.band_sensitivity,
                    expected_bands=exp_bands
                )
                st.session_state.band_detection_result = result
                st.session_state.detected_bands = result["bands"]
                st.success(f"在胶道 {st.session_state.marker_lane_idx+1} 检测到 {len(result['bands'])} 个条带（原始峰值 {result['n_detected_peaks']} 个）")
                band_viz = get_band_projection_viz(
                    result["projection"], result["bands"], result["peaks"], result["valleys"],
                    expected_bands=exp_bands
                )
                band_img = image.copy()
                draw = ImageDraw.Draw(band_img)
                x1, x2 = lane_box
                for j, (y1, y2, intensity) in enumerate(result["bands"]):
                    draw.rectangle([x1, y1, x2, y2], outline=(255, 0, 0), width=2)
                    draw.text((x1 + 2, y1), f"{j+1}", fill=(255, 0, 0))
                bcol1, bcol2 = st.columns([1, 1])
                with bcol1:
                    st.image(band_img, caption="条带覆盖图（红框=条带）", width=250)
                with bcol2:
                    st.image(cv2_to_pil(band_viz), caption="📊 垂直投影分析图 | 绿色=峰(条带中心) | 红色=谷 | 橙框=条带", width=250)
            else:
                st.error("请先选择有效的Marker胶道")

# ========== Tab 2: Marker校准 ==========
with tab2:
    st.subheader("Marker 选择与校准")
    
    all_markers = load_markers()
    preset_names = [k for k, v in all_markers.items() if v.get("source") == "preset"]
    custom_names = [k for k, v in all_markers.items() if v.get("source") == "custom"]
    
    col_m1, col_m2 = st.columns([1, 1])
    
    with col_m1:
        st.markdown("**选择预设Marker（来自 mk_standards.json）**")
        source_options = []
        if preset_names:
            source_options.append("📚 预设标准库")
        if custom_names:
            source_options.append("🔧 用户自定义")
        source_options.append("📤 上传标准图校准")
        
        selected_source = st.radio("数据来源", source_options, index=0)
        
        if selected_source == "📚 预设标准库":
            selected_marker = st.selectbox("预设Marker", preset_names)
            marker_data = all_markers[selected_marker]
            st.session_state.selected_marker_source = "preset"
            st.session_state.selected_marker_name = selected_marker
        elif selected_source == "🔧 用户自定义":
            selected_marker = st.selectbox("自定义Marker", custom_names)
            marker_data = all_markers[selected_marker]
            st.session_state.selected_marker_source = "preset"
            st.session_state.selected_marker_name = selected_marker
        else:
            st.session_state.selected_marker_source = "upload"
            marker_data = None
        
        if marker_data:
            st.markdown(f"**条带信息:** {marker_data.get('description', '')}")
            st.markdown(f"**分子量:** {marker_data['bands']}")
            unit = marker_data.get('unit', 'kDa')
            
            if len(st.session_state.detected_bands) > 0:
                mapping = map_bands_to_kda(
                    st.session_state.detected_bands,
                    marker_data['bands'],
                    st.session_state.gel_image.height if st.session_state.gel_image else 1000
                )
                st.session_state.band_kda_mapping = mapping
                
                st.markdown("**条带-kDa映射:**")
                mapped_bands = []
                for i, band in enumerate(st.session_state.detected_bands):
                    y_center = (band[0] + band[1]) / 2
                    kda = mapping.get(i, "?")
                    mapped_bands.append({
                        "y": y_center,
                        "kda": kda,
                        "unit": unit,
                        "box": (band[0], band[1])
                    })
                    st.markdown(f"- 条带 {i+1} (y≈{int(y_center)}): **{kda} {unit}**")
                #st.session_state.mapped_bands = mapped_bands
                # 只在首次映射或条带数量变化时自动赋值，避免覆盖用户在Tab3中的手动修改
                if not st.session_state.get("mapped_bands") or \
                    len(st.session_state.mapped_bands) != len(st.session_state.detected_bands):
                    st.session_state.mapped_bands = mapped_bands
    with col_m2:
        if selected_source == "📤 上传标准图校准":
            st.markdown("**上传Marker标准图**")
            ref_file = st.file_uploader("上传Marker标准图",
                                       type=["jpg", "jpeg", "png"],
                                       key="ref_uploader")
            
            if ref_file is not None:
                ref_img = Image.open(ref_file).convert("RGB")
                st.image(ref_img, caption="参考图", width=200)
                
                # OCR 识别图中数字
                with st.spinner("🔍 正在OCR识别图中数字..."):
                    ocr_numbers = ocr_marker_numbers(ref_img)
                
                if ocr_numbers:
                    st.success(f"识别到 {len(ocr_numbers)} 个数字（从上到下）: {ocr_numbers}")
                else:
                    st.error("未识别到有效数字，请检查图片清晰度")
                    ocr_numbers = []
                
                st.markdown("**从上到下核对/修改各条带kDa：**")
                ref_band_kdas = []
                for i, default_val in enumerate(ocr_numbers):
                    is_int = default_val == int(default_val)
                    step = 1 if is_int else 0.1
                    value = int(default_val) if is_int else float(default_val)
                    kda = st.number_input(
                        f"条带 {i+1}",
                        value=value,
                        step=step,
                        key=f"ref_kda_{i}"
                    )
                    ref_band_kdas.append(kda)
                
                # 如果OCR没识别到，允许手动添加
                if not ocr_numbers:
                    manual_count = st.number_input("手动输入条带数", 1, 30, 10)
                    ref_band_kdas = []
                    for i in range(manual_count):
                        kda = st.number_input(f"条带 {i+1}", value=10, step=1, key=f"manual_kda_{i}")
                        ref_band_kdas.append(kda)
                
                st.session_state.ref_band_kdas = ref_band_kdas
                
                if st.button("✅ 应用此kDa到当前Marker", type="primary"):
                    if len(st.session_state.detected_bands) > 0:
                        current_bands = st.session_state.detected_bands
                        n_current = len(current_bands)
                        n_ref = len(ref_band_kdas)
                        
                        if n_current == n_ref:
                            mapped = []
                            for i, band in enumerate(current_bands):
                                y_center = (band[0] + band[1]) / 2
                                mapped.append({
                                    "y": y_center,
                                    "kda": ref_band_kdas[i],
                                    "unit": "kDa",
                                    "box": (band[0], band[1])
                                })
                            st.session_state.mapped_bands = mapped
                            st.success(f"已应用 {n_current} 个kDa标注")
                        else:
                            mapped = []
                            for i, band in enumerate(current_bands):
                                y_center = (band[0] + band[1]) / 2
                                ref_idx = min(int(i * n_ref / n_current), n_ref - 1)
                                mapped.append({
                                    "y": y_center,
                                    "kda": ref_band_kdas[ref_idx],
                                    "unit": "kDa",
                                    "box": (band[0], band[1])
                                })
                            st.session_state.mapped_bands = mapped
                            st.warning(f"条带数不一致（检测{n_current} vs 标准{n_ref}），已按位置最近匹配")
                
                st.divider()
                st.markdown("**保存到标准库**")
                new_preset_name = st.text_input("新预设名称", value=f"自定义_{datetime.now().strftime('%m%d')}")
                if st.button("💾 保存到 mk_standards.json"):
                    save_preset_marker(new_preset_name, ref_band_kdas, new_preset_name)
                    st.success(f"已保存到标准库: {new_preset_name}")
                    st.rerun()
        else:
            st.info("选择「上传标准图校准」以添加新的Marker标准")

    st.divider()
    st.subheader("自定义Marker管理")
    custom_markers = {k: v for k, v in all_markers.items() if v.get("source") == "custom"}
    if custom_markers:
        for name, data in custom_markers.items():
            c1, c2 = st.columns([4, 1])
            with c1:
                st.markdown(f"**{name}**: {data['bands']} {data.get('unit', 'kDa')}")
            with c2:
                if st.button("🗑️ 删除", key=f"del_marker_{name}"):
                    from utils.gel_db import delete_custom_marker
                    delete_custom_marker(name)
                    st.success(f"已删除 {name}")
                    st.rerun()
    else:
        st.info("暂无自定义Marker")

# ========== Tab 3: 预览与保存 ==========
with tab3:
    st.subheader("预览、精修与保存")
    
    if st.session_state.gel_image is not None and len(st.session_state.lane_boxes) > 0:
        
        left_col, right_col = st.columns([1, 2])
        with left_col:
            st.subheader("精修偏移")
            
            st.markdown("**📝 胶道名称精修**")
            lane_options_fine = list(range(len(st.session_state.lane_boxes)))
            if lane_options_fine:
                selected_lane_for_fine = st.selectbox(
                    "选择胶道",
                    lane_options_fine,
                    format_func=lambda i: f"胶道 {i+1}: {st.session_state.lane_names[i][:10]}",
                    key="lane_fine_sel"
                )
                lane_fine_idx = selected_lane_for_fine
                
                # 胶道名称编辑：key加索引，首次创建时从 lane_names 初始化
                # 同步 lane_names 到 Tab3 编辑 widget keys（必须在 widget 实例化前执行）
                for i in range(len(st.session_state.lane_names)):
                    key = f"edit_lane_name_{i}"
                    if key in st.session_state:
                        st.session_state[key] = st.session_state.lane_names[i]
                
                # 胶道名称编辑：key加索引，首次创建时从 lane_names 初始化
                key_name = f"edit_lane_name_{lane_fine_idx}"
                if key_name not in st.session_state:
                    st.session_state[key_name] = st.session_state.lane_names[lane_fine_idx]
                new_lane_name = st.text_input("胶道名称", key=key_name)

                # 必须复制整个列表再赋值，否则 Streamlit 检测不到列表内部变化
                lane_names_copy = list(st.session_state.lane_names)
                lane_names_copy[lane_fine_idx] = new_lane_name
                st.session_state.lane_names = lane_names_copy
                # 胶道偏移编辑
                current_lane_fine = st.session_state.lane_fine_offsets.get(lane_fine_idx, {"x": 0, "y": 0})
                lane_fx = st.number_input("胶道 X偏移", value=current_lane_fine.get("x", 0), step=1, key=f"lane_fx_{lane_fine_idx}")
                lane_fy = st.number_input("胶道 Y偏移", value=current_lane_fine.get("y", 0), step=1, key=f"lane_fy_{lane_fine_idx}")
                lane_fine_copy = dict(st.session_state.lane_fine_offsets)
                lane_fine_copy[lane_fine_idx] = {"x": lane_fx, "y": lane_fy}
                st.session_state.lane_fine_offsets = lane_fine_copy
            
            st.divider()
            st.markdown("**🧬 MK条带kDa精修**")
            mapped = getattr(st.session_state, 'mapped_bands', [])
            if mapped:
                band_options = list(range(len(mapped)))
                selected_band_for_fine = st.selectbox(
                    "选择条带",
                    band_options,
                    format_func=lambda i: f"条带 {i+1}: {mapped[i]['kda']}kDa (y≈{int(mapped[i]['y'])})",
                    key="band_fine_sel"
                )
                band_fine_idx = selected_band_for_fine

                # kDa 数值编辑
                current_kda = mapped[band_fine_idx]['kda']
                is_int_kda = isinstance(current_kda, int) or (isinstance(current_kda, float) and current_kda == int(current_kda))
                new_kda = st.number_input(
                    "kDa 数值",
                    value=float(current_kda) if not is_int_kda else int(current_kda),
                    step=0.1 if not is_int_kda else 1,
                    key=f"edit_kda_{band_fine_idx}"
                )
                # 必须深拷贝整个列表再赋值
                mapped_copy = [dict(b) for b in mapped]
                mapped_copy[band_fine_idx]['kda'] = new_kda
                st.session_state.mapped_bands = mapped_copy
                
                # 条带偏移编辑
                current_band_fine = st.session_state.band_fine_offsets.get(band_fine_idx, {"x": 0, "y": 0})
                band_fx = st.number_input("条带 X偏移", value=current_band_fine.get("x", 0), step=1, key=f"band_fx_{band_fine_idx}")
                band_fy = st.number_input("条带 Y偏移", value=current_band_fine.get("y", 0), step=1, key=f"band_fy_{band_fine_idx}")
                band_fine_copy = dict(st.session_state.band_fine_offsets)
                band_fine_copy[band_fine_idx] = {"x": band_fx, "y": band_fy}
                st.session_state.band_fine_offsets = band_fine_copy
            else:
                st.info("请先检测并校准Marker条带")
        with right_col:
            st.subheader("标注预览")
            
            if show_preview:
                mapped = getattr(st.session_state, 'mapped_bands', None)
                
                annotated = create_annotated_image(
                    original_image=st.session_state.gel_image,
                    lane_names=st.session_state.lane_names,
                    lane_boxes=st.session_state.lane_boxes,
                    marker_lane_idx=st.session_state.marker_lane_idx,
                    marker_type=getattr(st.session_state, 'selected_marker_name', None),
                    marker_bands_data=mapped,
                    font_name=selected_font,
                    lane_label_y_offset=lane_label_y_offset,
                    lane_label_x_offset=lane_label_x_offset,
                    lane_label_angle=lane_label_angle,
                    kda_label_x_offset=kda_label_x_offset,
                    kda_label_y_offset=kda_label_y_offset,
                    lane_fine_offsets=st.session_state.lane_fine_offsets,
                    band_fine_offsets=st.session_state.band_fine_offsets,
                    show_lane_boxes=show_lane_boxes
                )
                
                st.image(annotated, caption="标注预览（顶部和右侧已扩展白边）", width='stretch')
                st.session_state.last_annotated = annotated
        
        st.divider()
        col_s1, col_s2 = st.columns([2, 1])
        
        with col_s1:
            default_name = st.session_state.annotation_name or get_next_default_name()
            save_name = st.text_input("图片命名", value=default_name,
                                    placeholder="不填则自动命名为 p1, p2...")
            st.session_state.annotation_name = save_name
        
        with col_s2:
            st.markdown("&nbsp;")
            if st.button("💾 保存标注", type="primary", use_container_width=True):
                if hasattr(st.session_state, 'last_annotated') and st.session_state.last_annotated:
                    img_path = os.path.join(GEL_IMAGES_DIR, f"{save_name}.png")
                    original_path = os.path.join(GEL_IMAGES_DIR, f"{save_name}_original.png")
                    
                    counter = 1
                    base_name = save_name
                    while os.path.exists(img_path):
                        save_name = f"{base_name}_{counter}"
                        img_path = os.path.join(GEL_IMAGES_DIR, f"{save_name}.png")
                        original_path = os.path.join(GEL_IMAGES_DIR, f"{save_name}_original.png")
                        counter += 1
                    
                    st.session_state.last_annotated.save(img_path, "PNG")
                    if st.session_state.gel_image:
                        st.session_state.gel_image.save(original_path, "PNG")
                    # 同时保存原始未标注图，供后续编辑使用
                    if st.session_state.gel_image:
                        st.session_state.gel_image.save(original_path, "PNG")

                    annotation = {
                        "id": save_name,
                        "name": save_name,
                        "created_at": datetime.now().isoformat(),
                        "updated_at": datetime.now().isoformat(),
                        "image_path": img_path,
                        "original_image_path": original_path,
                        "lane_names": st.session_state.lane_names,
                        "lane_boxes": st.session_state.lane_boxes,
                        "marker_lane_idx": st.session_state.marker_lane_idx,
                        "marker_type": getattr(st.session_state, 'selected_marker_name', None),
                        "marker_bands": getattr(st.session_state, 'mapped_bands', []),
                        "font_name": selected_font,
                        "lane_label_y_offset": lane_label_y_offset,
                        "lane_label_x_offset": lane_label_x_offset,
                        "lane_label_angle": lane_label_angle,
                        "kda_label_x_offset": kda_label_x_offset,
                        "kda_label_y_offset": kda_label_y_offset,
                        "lane_fine_offsets": st.session_state.lane_fine_offsets,
                        "band_fine_offsets": st.session_state.band_fine_offsets,
                        "lane_individual_shifts": st.session_state.lane_individual_shifts,
                        "global_x_shift": st.session_state.global_x_shift,
                        "selected_lane_method": st.session_state.selected_lane_method,
                    }
                    
                    save_annotation(annotation)
                    st.success(f"✅ 已保存: {save_name}")
                    st.balloons()
                else:
                    st.error("请先生成预览")
            st.markdown("&nbsp;")
            if st.button("📊 生成 PPTX", use_container_width=True):
                with st.spinner("正在生成 PPTX..."):
                    pptx_buf = create_annotated_pptx(
                        original_image=st.session_state.gel_image,
                        lane_names=st.session_state.lane_names,
                        lane_boxes=st.session_state.lane_boxes,
                        marker_lane_idx=st.session_state.marker_lane_idx,
                        marker_bands_data=getattr(st.session_state, 'mapped_bands', None),
                        font_name=selected_font,
                        lane_label_y_offset=lane_label_y_offset,
                        lane_label_x_offset=lane_label_x_offset,
                        lane_label_angle=lane_label_angle,
                        kda_label_x_offset=kda_label_x_offset,
                        kda_label_y_offset=kda_label_y_offset,
                        lane_fine_offsets=st.session_state.lane_fine_offsets,
                        band_fine_offsets=st.session_state.band_fine_offsets,
                    )
                    st.session_state.pptx_buffer = pptx_buf.getvalue()
                    st.session_state.pptx_name = save_name
                st.rerun()

            if st.session_state.get("pptx_buffer"):
                st.download_button(
                    label="⬇️ 下载 PPTX",
                    data=st.session_state["pptx_buffer"],
                    file_name=f"{st.session_state.get('pptx_name', 'gel')}.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True,
                )
    else:
        st.info("请先上传图像并完成胶道识别")

st.divider()
st.caption("💡 提示：保存后的标注可以在「胶图标注管理」页面进行人工修改")

# ========== Tab 4: 标注管理 ==========
with tab4:
    st.subheader("📝 已保存的标注项目管理")

    from utils.gel_db import load_annotations, get_annotation_by_id, delete_annotation

    annotations = load_annotations()

    if len(annotations) == 0:
        st.info("📭 暂无保存的标注项目")
    else:
        # 搜索
        search = st.text_input("🔍 搜索项目名称", placeholder="输入名称过滤...")
        filtered = [a for a in annotations if search.lower() in a.get("name", "").lower()] if search else annotations

        if not filtered:
            st.warning("未找到匹配的项目")
        else:
            # 网格展示
            cols = st.columns(3)
            for i, ann in enumerate(filtered):
                with cols[i % 3]:
                    with st.container(border=True):
                        img_path = ann.get("image_path", "")
                        if os.path.exists(img_path):
                            st.image(img_path, width='stretch')
                        else:
                            st.error("图片丢失")

                        st.markdown(f"**{ann.get('name', '未命名')}**")
                        st.caption(f"创建: {ann.get('created_at', '')[:10]}")

                        c1, c2 = st.columns(2)
                        with c1:
                            if st.button("✏️ 编辑", key=f"edit_{ann['id']}", use_container_width=True):
                                st.session_state.editing_id = ann["id"]
                                st.rerun()
                        with c2:
                            if st.button("🗑️ 删除", key=f"del_{ann['id']}", use_container_width=True):
                                st.session_state.confirm_delete = ann["id"]
                                st.rerun()

                        if st.session_state.get("confirm_delete") == ann["id"]:
                            st.warning("⚠️ 确定删除？不可撤销！")
                            d1, d2 = st.columns(2)
                            with d1:
                                if st.button("✅ 确认", key=f"confirm_del_{ann['id']}", type="primary", use_container_width=True):
                                    delete_annotation(ann["id"])
                                    st.session_state.confirm_delete = None
                                    st.success("已删除")
                                    st.rerun()
                            with d2:
                                if st.button("❌ 取消", key=f"cancel_del_{ann['id']}", use_container_width=True):
                                    st.session_state.confirm_delete = None
                                    st.rerun()

                        # 下载
                        if os.path.exists(img_path):
                            with open(img_path, "rb") as f:
                                st.download_button("⬇️ 下载", f.read(),
                                                  file_name=f"{ann.get('name', 'gel')}_annotated.png",
                                                  mime="image/png",
                                                  use_container_width=True)

    # 编辑模式（简单版：加载到当前session重新编辑）
    if st.session_state.get("editing_id") and not st.session_state.get("editing_loaded"):
        ann = get_annotation_by_id(st.session_state.editing_id)
        if ann:
            st.divider()
            st.subheader(f"✏️ 编辑: {ann.get('name', '')}")
            if st.button("⬅️ 返回列表"):
                st.session_state.editing_id = None
                st.session_state.editing_loaded = False
                st.rerun()

            # 加载数据到当前session（重新检测太复杂，直接加载已有标注数据）
            st.info("编辑功能：加载已有数据到当前标注流程。请切换到「预览与保存」标签查看和修改。")

            # JSON序列化后字典key会变成字符串，需递归转回整数
            def int_keys(d):
                if not isinstance(d, dict):
                    return d
                return {int(k) if str(k).lstrip('-').isdigit() else k: int_keys(v) for k, v in d.items()}

            st.session_state.lane_names = ann.get("lane_names", [])
            st.session_state.lane_boxes = [tuple(b) for b in ann.get("lane_boxes", [])]
            st.session_state.marker_lane_idx = ann.get("marker_lane_idx")
            st.session_state.mapped_bands = ann.get("marker_bands", [])
            st.session_state.lane_fine_offsets = int_keys(ann.get("lane_fine_offsets", {}))
            st.session_state.band_fine_offsets = int_keys(ann.get("band_fine_offsets", {}))
            st.session_state.lane_individual_shifts = int_keys(ann.get("lane_individual_shifts", {}))
            st.session_state.global_x_shift = ann.get("global_x_shift", 0)
            st.session_state.selected_lane_method = ann.get("selected_lane_method", "A")
            
            # 恢复方法A/B的基础数据
            if st.session_state.selected_lane_method == "A":
                st.session_state.method_a_lanes = [tuple(b) for b in ann.get("lane_boxes", [])]
                st.session_state.method_a_expected = len(ann.get("lane_boxes", []))
            else:
                st.session_state.method_b_lanes = [tuple(b) for b in ann.get("lane_boxes", [])]
                st.session_state.method_b_all_peaks = [(b[0] + b[1]) // 2 for b in ann.get("lane_boxes", [])]
                st.session_state.method_b_expected = len(ann.get("lane_boxes", []))

            # 加载图片：优先加载原始未标注图，避免重复叠加文字
            if os.path.exists(ann.get("original_image_path", "")):
                st.session_state.gel_image = Image.open(ann["original_image_path"]).convert("RGB")
                st.success("数据已加载，请切换到「预览与保存」标签修改")
            elif os.path.exists(ann.get("image_path", "")):
                st.session_state.gel_image = Image.open(ann["image_path"]).convert("RGB")
                st.warning("⚠️ 未找到原始图，使用标注图编辑（可能出现重复标注）")
            else:
                st.error("原始图片已丢失，无法编辑")
            
            # 标记已加载，后续 rerun 不再重复覆盖
            st.session_state.editing_loaded = True
        else:
            st.error("项目不存在")
            st.session_state.editing_id = None
            st.session_state.editing_loaded = False
