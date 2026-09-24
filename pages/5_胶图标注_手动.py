import streamlit as st
import numpy as np
from PIL import Image, ImageDraw, ImageFont
from streamlit_paste_button import paste_image_button as pbutton
import os
import math
from utils.storage import (
    load_mk_standards, save_project, list_projects,
    load_project, delete_project, add_mk_standard, reset_mk_standards
)
from streamlit_image_coordinates import streamlit_image_coordinates
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.dml.color import RGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False
    Presentation = None

st.title("🧬 胶图标注工具")
st.sidebar.title("功能选择")
subpage = st.sidebar.radio("选择操作", ["手工标注", "管理已保存项目"])

def draw_markers(image_pil, lanes):
    """绘制红点(胶道)和蓝点(条带)标记"""
    img = image_pil.copy()
    draw = ImageDraw.Draw(img)
    for lane in lanes:
        x = lane['x']
        draw.ellipse((x - 5, 0, x + 5, 10), fill='red', outline='red')
    for lane in lanes:
        x = lane['x']
        for y in lane.get('bands', []):
            draw.ellipse((x - 4, y - 4, x + 4, y + 4), fill='blue', outline='blue')
    return img


def load_font(font_name, font_size):
    """加载系统TrueType字体，确保抗锯齿和字号生效"""
    possible_names = [
        font_name,
        font_name.lower(),
        font_name.capitalize(),
        f"{font_name}.ttf",
        f"{font_name}.TTF",
    ]
    search_dirs = [
        "/System/Library/Fonts/",
        "C:/Windows/Fonts/",
        "/usr/share/fonts/truetype/",
        "/usr/local/share/fonts/",
        os.path.expanduser("~/.fonts/"),
        os.path.expanduser("~/Library/Fonts/"),
    ]
    aliases = {
        "Arial": ["Arial.ttf", "arial.ttf", "LiberationSans-Regular.ttf", "DejaVuSans.ttf"],
        "Times New Roman": ["TimesNewRoman.ttf", "times.ttf", "LiberationSerif-Regular.ttf"],
        "Courier New": ["CourierNew.ttf", "cour.ttf", "LiberationMono-Regular.ttf"],
        "SimHei": ["SimHei.ttf", "simhei.ttf", "NotoSansCJK-Regular.ttc"],
    }
    candidates = aliases.get(font_name, []) + possible_names
    for name in candidates:
        for base in search_dirs:
            path = os.path.join(base, name)
            if os.path.exists(path):
                try:
                    return ImageFont.truetype(path, font_size)
                except:
                    continue
    try:
        return ImageFont.truetype(font_name, font_size)
    except:
        pass
    st.warning(f"⚠️ 未找到字体 '{font_name}'，使用默认位图字体（清晰度较低）")
    return ImageFont.load_default()


def draw_text_with_rotation(draw, xy, text, fill, font, angle=0, anchor='lt'):
    """绘制旋转文本，确保无裁剪"""
    bbox = draw.textbbox((0, 0), text, font=font, anchor='lt')
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    pad = 30
    temp = Image.new('RGBA', (tw + 2 * pad, th + 2 * pad), (0, 0, 0, 0))
    temp_draw = ImageDraw.Draw(temp)
    temp_draw.text((pad, pad), text, font=font, fill=fill, anchor='lt')
    rotated = temp.rotate(angle, expand=True, resample=Image.BICUBIC, fillcolor=(0, 0, 0, 0))
    rw, rh = rotated.size

    anchor_offsets = {
        'lt': (0, 0),
        'mt': (tw/2, 0),
        'rt': (tw, 0),
        'lm': (0, th/2),
        'mm': (tw/2, th/2),
        'rm': (tw, th/2),
        'lb': (0, th),
        'mb': (tw/2, th),
        'rb': (tw, th)
    }
    ox, oy = anchor_offsets.get(anchor, (0, 0))
    ox += pad
    oy += pad

    cx = (tw + 2 * pad) / 2
    cy = (th + 2 * pad) / 2
    rad = math.radians(angle)
    dx = ox - cx
    dy = oy - cy
    new_dx = dx * math.cos(rad) - dy * math.sin(rad)
    new_dy = dx * math.sin(rad) + dy * math.cos(rad)
    new_ox = cx + new_dx
    new_oy = cy + new_dy

    paste_x = int(xy[0] - new_ox)
    paste_y = int(xy[1] - new_oy)
    draw._image.paste(rotated, (paste_x, paste_y), rotated)

def generate_svg_overlay(lanes, mk_standards, chosen_mk,
                         label_offset, rotation, font_name, band_offset,
                         font_size, lane_color, band_color,
                         offset_dict, img_width, img_height, background_img):
    """
    生成 SVG 标注图层，背景为原图（可编辑文字）
    background_img: PIL Image 对象（裁剪后的胶图）
    """
    import xml.etree.ElementTree as ET
    from xml.dom import minidom
    import base64
    from io import BytesIO

    # 将背景图转为 base64
    buffered = BytesIO()
    background_img.save(buffered, format="PNG")
    img_base64 = base64.b64encode(buffered.getvalue()).decode()

    # SVG 根节点
    svg = ET.Element('svg', xmlns="http://www.w3.org/2000/svg",
                     width=str(img_width), height=str(img_height),
                     viewBox=f"0 0 {img_width} {img_height}")

    # 嵌入背景图片
    image_elem = ET.SubElement(svg, 'image',
                               href=f"data:image/png;base64,{img_base64}",
                               width=str(img_width), height=str(img_height))

    font_family = font_name if font_name in ["Arial", "Times New Roman", "Courier New", "SimHei"] else "sans-serif"

    # 收集所有标注项
    text_items = []
    for idx, lane in enumerate(lanes):
        x = lane['x']
        name = lane['name']
        text_items.append({
            'x': x,
            'y': label_offset,
            'text': name,
            'color': lane_color,
            'anchor': 'mt',
            'rotate': rotation,
            'key': (idx, 'label')
        })
        if name.upper() == "MK" and chosen_mk in mk_standards:
            kda_list = mk_standards[chosen_mk]
            bands = sorted(lane.get('bands', []))
            for bi, y in enumerate(bands):
                if bi < len(kda_list):
                    text_items.append({
                        'x': x - band_offset,
                        'y': y,
                        'text': str(kda_list[bi]),
                        'color': band_color,
                        'anchor': 'rm',
                        'rotate': 0,
                        'key': (idx, 'band', bi)
                    })

    # 应用精细偏移
    for item in text_items:
        key = item['key']
        if key in offset_dict:
            dx, dy = offset_dict[key]
            item['x'] += dx
            item['y'] += dy

    # 调整顶部偏移（如果文字超出画布，移动整体）
    min_y_all = 0
    for item in text_items:
        approx_height = font_size * 1.2
        if item['anchor'] in ('mt', 'rm'):
            min_y = item['y'] - approx_height / 2
        else:
            min_y = item['y']
        if min_y < min_y_all:
            min_y_all = min_y
    if min_y_all < 0:
        shift_y = -min_y_all + 10
        for item in text_items:
            item['y'] += shift_y
        svg.set('height', str(img_height + shift_y))
        svg.set('viewBox', f"0 0 {img_width} {img_height + shift_y}")
        # 同时调整背景图片高度
        image_elem.set('height', str(img_height + shift_y))

    # 添加文本元素
    # 添加文本元素
    anchor_map = {
        'lt': ('start', 'baseline'),
        'mt': ('middle', 'baseline'),
        'rt': ('end', 'baseline'),
        'lm': ('start', 'middle'),
        'mm': ('middle', 'middle'),
        'rm': ('end', 'middle'),
        'lb': ('start', 'hanging'),
        'mb': ('middle', 'hanging'),
        'rb': ('end', 'hanging')
    }
    for item in text_items:
        text_anchor, dominant_baseline = anchor_map.get(item['anchor'], ('start', 'baseline'))
        # 修正旋转方向：PIL 正角度逆时针，SVG 正角度顺时针，取反
        transform = f"rotate({-item['rotate']}, {item['x']}, {item['y']})" if item['rotate'] != 0 else ""
        color = item['color'] if item['color'].startswith('#') else '#000000'
        text_elem = ET.SubElement(svg, 'text',
                                  x=str(item['x']),
                                  y=str(item['y']),
                                  fill=color,
                                  font_size=str(font_size),
                                  font_family=font_family,
                                  text_anchor=text_anchor,
                                  dominant_baseline=dominant_baseline,
                                  transform=transform)
        text_elem.text = item['text']

    rough_string = ET.tostring(svg, 'utf-8')
    reparsed = minidom.parseString(rough_string)
    return reparsed.toprettyxml(indent="  ")

def generate_pptx(background_img, lanes, mk_standards, chosen_mk,
                  label_offset, rotation, font_name, band_offset,
                  font_size, lane_color, band_color, offset_dict):
    """
    生成 PPTX 文件（字节流），包含背景图和可编辑的标注文本框。
    文字旋转角度与预览一致，位置通过几何补偿精确对齐。
    """
    if not PPTX_AVAILABLE:
        raise RuntimeError("请先安装 python-pptx：pip install python-pptx")

    from io import BytesIO
    import math

    img_width, img_height = background_img.size

    # 创建演示文稿，尺寸与图片比例一致（96 DPI）
    prs = Presentation()
    prs.slide_width = Inches(img_width / 96.0)
    prs.slide_height = Inches(img_height / 96.0)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 背景图片
    img_stream = BytesIO()
    background_img.save(img_stream, format='PNG')
    img_stream.seek(0)
    slide.shapes.add_picture(img_stream, 0, 0,
                             width=prs.slide_width, height=prs.slide_height)

    # 加载字体（用于计算文本尺寸）
    font = load_font(font_name, font_size)

    # 构建文本项（同 create_annotation_image）
    text_items = []
    for idx, lane in enumerate(lanes):
        x = lane['x']
        name = lane['name']
        text_items.append({
            'x': x,
            'y': label_offset,
            'text': name,
            'color': lane_color,
            'anchor': 'mt',
            'rotate': rotation,
            'key': (idx, 'label')
        })
        if name.upper() == "MK" and chosen_mk in mk_standards:
            kda_list = mk_standards[chosen_mk]
            bands = sorted(lane.get('bands', []))
            for bi, y in enumerate(bands):
                if bi < len(kda_list):
                    text_items.append({
                        'x': x - band_offset,
                        'y': y,
                        'text': str(kda_list[bi]),
                        'color': band_color,
                        'anchor': 'rm',
                        'rotate': 0,          # kDa 不旋转
                        'key': (idx, 'band', bi)
                    })

    # 应用精细偏移
    for item in text_items:
        key = item['key']
        if key in offset_dict:
            dx, dy = offset_dict[key]
            item['x'] += dx
            item['y'] += dy

    # 用于计算文本尺寸的临时绘图
    from PIL import Image, ImageDraw
    temp_img = Image.new('RGB', (1, 1))
    draw = ImageDraw.Draw(temp_img)

    # 锚点相对于文本框左上角的偏移比例（像素）
    anchor_map = {
        'lt': (0, 0),
        'mt': (0.5, 0),
        'rt': (1, 0),
        'lm': (0, 0.5),
        'mm': (0.5, 0.5),
        'rm': (1, 0.5),
        'lb': (0, 1),
        'mb': (0.5, 1),
        'rb': (1, 1)
    }

    for item in text_items:
        # 获取文本的准确尺寸
        bbox = draw.textbbox((0, 0), item['text'], font=font)
        tw = bbox[2] - bbox[0]
        th = bbox[3] - bbox[1]

        # 锚点相对文本框左上角的偏移（像素）
        anchor_rel = anchor_map.get(item['anchor'], (0.5, 0.5))
        ax = tw * anchor_rel[0]
        ay = th * anchor_rel[1]

        # 目标锚点位置（像素坐标）
        target_x = item['x']
        target_y = item['y']

        # 旋转角度（PPT 顺时针为正，PIL 逆时针为正，取反）
        angle_deg = -item['rotate'] if item['rotate'] != 0 else 0
        angle_rad = math.radians(angle_deg)

        # 若旋转角度不为0，需补偿文本框位置
        if abs(angle_deg) > 0.1:
            # 锚点相对文本框中心的偏移
            cx_offset = ax - tw / 2.0
            cy_offset = ay - th / 2.0

            # 旋转后偏移
            rot_cx = cx_offset * math.cos(angle_rad) - cy_offset * math.sin(angle_rad)
            rot_cy = cx_offset * math.sin(angle_rad) + cy_offset * math.cos(angle_rad)

            # 文本框中心应位于：目标锚点 - 旋转后的偏移
            center_x = target_x - rot_cx
            center_y = target_y - rot_cy

            # 左上角坐标
            left_px = center_x - tw / 2.0
            top_px = center_y - th / 2.0
        else:
            # 无旋转，直接根据锚点计算左上角
            left_px = target_x - ax
            top_px = target_y - ay

        # 转换为英寸（96 DPI）
        left_in = Inches(left_px / 96.0)
        top_in = Inches(top_px / 96.0)
        width_in = Inches(tw / 96.0 + 0.05)   # 加极小边距防止文字溢出
        height_in = Inches(th / 96.0 + 0.05)

        # 添加文本框
        textbox = slide.shapes.add_textbox(left_in, top_in, width_in, height_in)
        # 设置旋转
        if abs(angle_deg) > 0.1:
            textbox.rotation = angle_deg

        # 设置文本内容和样式
        tf = textbox.text_frame
        tf.text = item['text']
        p = tf.paragraphs[0]
        p.font.name = font_name if font_name in ["Arial", "Times New Roman", "Courier New", "SimHei"] else "Arial"
        p.font.size = Pt(font_size * 0.75)   # 像素转磅近似
        # 颜色解析
        color_hex = item['color'].lstrip('#')
        r, g, b = tuple(int(color_hex[i:i+2], 16) for i in (0, 2, 4))
        p.font.color.rgb = RGBColor(r, g, b)

        # 对齐方式（根据锚点设置文本框内文字对齐）
        # 由于我们已经通过位置补偿，文本框内文字无需再调整对齐
        # 但为了保险，设置段落对齐为居中或左对齐
        if item['anchor'] in ('mt', 'mm', 'mb'):
            p.alignment = PP_ALIGN.CENTER
        elif item['anchor'] in ('rt', 'rm', 'rb'):
            p.alignment = PP_ALIGN.RIGHT
        else:
            p.alignment = PP_ALIGN.LEFT

        # 垂直对齐（文本在文本框内垂直居中）
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    # 保存为字节流
    pptx_io = BytesIO()
    prs.save(pptx_io)
    pptx_io.seek(0)
    return pptx_io.getvalue()

def create_annotation_image(image_pil, lanes, mk_standards, chosen_mk,
                            label_offset, rotation, font_name, band_offset,
                            font_size, lane_color, band_color,
                            offset_dict=None):
    """
    绘制标注，支持全局偏移和精细偏移（offset_dict）。
    offset_dict键: (lane_idx, 'label') 或 (lane_idx, 'band', band_idx)
    注意：MK条带会根据y坐标升序自动排序后匹配kDa列表。
    """
    if offset_dict is None:
        offset_dict = {}
    img = image_pil.copy()
    draw = ImageDraw.Draw(img)
    font = load_font(font_name, font_size)

    text_items = []
    for idx, lane in enumerate(lanes):
        x = lane['x']
        name = lane['name']
        # 胶道名称
        text_items.append({
            'x': x,
            'y': label_offset,
            'text': name,
            'color': lane_color,
            'anchor': 'mt',
            'rotate': rotation,
            'key': (idx, 'label')
        })
        # kDa标注（仅MK且有效）
        if name.upper() == "MK" and chosen_mk in mk_standards:
            kda_list = mk_standards[chosen_mk]
            # 按y坐标升序排序（从上到下对应大分子量到小分子量）
            bands = sorted(lane.get('bands', []))
            for bi, y in enumerate(bands):
                if bi < len(kda_list):
                    text_items.append({
                        'x': x - band_offset,
                        'y': y,
                        'text': str(kda_list[bi]),
                        'color': band_color,
                        'anchor': 'rm',
                        'rotate': 0,
                        'key': (idx, 'band', bi)
                    })

    # 应用精细偏移
    for item in text_items:
        key = item['key']
        if key in offset_dict:
            dx, dy = offset_dict[key]
            item['x'] += dx
            item['y'] += dy

    # 计算扩展区域（顶部和右侧）
    top_pad = 0
    right_pad = 0
    img_w, img_h = img.size
    for item in text_items:
        bbox = draw.textbbox((0, 0), item['text'], font=font)
        tw = bbox[2] - bbox[0]
        th = bbox[3] - bbox[1]
        angle = math.radians(item['rotate'])
        rot_w = int(tw * abs(math.cos(angle)) + th * abs(math.sin(angle))) + 20
        rot_h = int(th * abs(math.cos(angle)) + tw * abs(math.sin(angle))) + 20
        if item['anchor'] in ('mt', 'rm'):
            min_y = item['y'] - rot_h // 2
        else:
            min_y = item['y']
        if min_y < 0:
            top_pad = max(top_pad, -min_y + 30)
        if item['anchor'] in ('rt', 'rm', 'rb'):
            max_x = item['x'] + rot_w
        elif item['anchor'] in ('mt', 'mm', 'mb'):
            max_x = item['x'] + rot_w // 2
        else:
            max_x = item['x']
        if max_x > img_w:
            right_pad = max(right_pad, max_x - img_w + 30)

    if top_pad > 0 or right_pad > 0:
        new_w = img_w + right_pad
        new_h = img_h + top_pad
        new_img = Image.new('RGB', (new_w, new_h), color='white')
        new_img.paste(img, (0, top_pad))
        img = new_img
        draw = ImageDraw.Draw(img)
        for item in text_items:
            item['y'] += top_pad

    for item in text_items:
        draw_text_with_rotation(
            draw,
            (item['x'], item['y']),
            item['text'],
            fill=item['color'],
            font=font,
            angle=item['rotate'],
            anchor=item['anchor']
        )

    return img

def show_annotation_interface(pil_img):
    """显示标注界面（裁剪、标记、预览、保存等）"""
    import numpy as np
    from streamlit_image_coordinates import streamlit_image_coordinates

    img_array = np.array(pil_img)

    # ---- 裁剪控制 ----
    with st.expander("✂️ 裁剪图像（去除四周空白）", expanded=False):
        col1, col2 = st.columns(2)
        with col1:
            left = st.slider("左边界裁剪 (像素)", 0, st.session_state.original_uploaded.width-1,
                             st.session_state.crop_left, key="crop_left_slider")
            top = st.slider("上边界裁剪 (像素)", 0, st.session_state.original_uploaded.height-1,
                            st.session_state.crop_top, key="crop_top_slider")
        with col2:
            right = st.slider("右边界裁剪 (像素)", 0, st.session_state.original_uploaded.width-1,
                              st.session_state.crop_right, key="crop_right_slider")
            bottom = st.slider("下边界裁剪 (像素)", 0, st.session_state.original_uploaded.height-1,
                               st.session_state.crop_bottom, key="crop_bottom_slider")
        if (left, top, right, bottom) != (st.session_state.crop_left, st.session_state.crop_top,
                                           st.session_state.crop_right, st.session_state.crop_bottom):
            st.session_state.crop_left = left
            st.session_state.crop_top = top
            st.session_state.crop_right = right
            st.session_state.crop_bottom = bottom
            new_cropped = apply_crop(st.session_state.original_uploaded, st.session_state)
            st.session_state.uploaded_img = new_cropped
            st.session_state.lanes = []
            st.session_state.preview_image = None
            st.rerun()

        col_orig, col_crop = st.columns(2)
        with col_orig:
            st.image(st.session_state.original_uploaded, caption="原始图", width='stretch')
        with col_crop:
            st.image(st.session_state.uploaded_img, caption="裁剪后", width='stretch')

    # ---- 模式提示 ----
    if st.session_state.mode == 'lane':
        st.info("💡 点击左图添加胶道中心点（红色圆点）")
    else:
        st.info("🔵 点击左图添加条带位置（蓝色圆点），系统自动按从上到下顺序匹配")

    # ---- 双列布局（点击 + 标记预览） ----
    col_click, col_preview = st.columns(2)
    with col_click:
        st.write("**点击图像**")
        value = streamlit_image_coordinates(img_array, key="image_click")
    with col_preview:
        st.write("**标记预览** (红点=胶道, 蓝点=条带)")
        preview_img = draw_markers(pil_img, st.session_state.lanes)
        st.image(preview_img, width='stretch')

    # ---- 处理点击事件 ----
    if value is not None:
        x, y = value["x"], value["y"]
        mk_standards = load_mk_standards()
        if st.session_state.mode == 'lane':
            if not any(abs(lane['x'] - x) < 5 for lane in st.session_state.lanes):
                st.session_state.lanes.append({
                    "x": x,
                    "name": f"Lane{len(st.session_state.lanes)+1}",
                    "bands": []
                })
                st.success(f"✅ 添加胶道 x={x}")
                st.rerun()
        elif st.session_state.mode == 'band':
            mk_idx = st.session_state.current_mk_idx
            if mk_idx == -1:
                st.error("未选择 MK 胶道")
            elif st.session_state.selected_mk_std not in mk_standards:
                st.error("请先选择有效的 MK 标准")
            else:
                existing = st.session_state.lanes[mk_idx]['bands']
                if not any(abs(y0 - y) < 3 for y0 in existing):
                    st.session_state.lanes[mk_idx]['bands'].append(y)
                    st.session_state.lanes[mk_idx]['bands'] = sorted(st.session_state.lanes[mk_idx]['bands'])
                    st.success(f"🔵 添加条带 y={y}")
                    st.rerun()
                else:
                    st.warning("该位置已有条带")

    # ---- 胶道列表 ----
    st.subheader("已标记的胶道")
    if st.session_state.lanes:
        for i, lane in enumerate(st.session_state.lanes):
            col1, col2, col3, col4 = st.columns([2, 2, 1, 1])
            with col1:
                st.write(f"胶道 {i+1}: x={lane['x']}")
            with col2:
                new_name = st.text_input("名称", value=lane['name'], key=f"lname_{i}")
                if new_name != lane['name']:
                    st.session_state.lanes[i]['name'] = new_name
                    if new_name.upper() == "MK":
                        st.session_state.lanes[i]['bands'] = []
                        if st.session_state.current_mk_idx == i:
                            st.session_state.current_mk_idx = -1
                    st.rerun()
            with col3:
                if st.button("删除", key=f"del_{i}"):
                    if st.session_state.current_mk_idx == i:
                        st.session_state.current_mk_idx = -1
                    del st.session_state.lanes[i]
                    st.rerun()
            with col4:
                st.write(f"条带: {len(lane['bands'])}")
    else:
        st.info("尚未标记胶道，请点击左图添加")

    # ---- MK 条带标记 ----
    mk_standards = load_mk_standards()
    mk_options = list(mk_standards.keys()) + ["自定义..."]
    mk_indices = [i for i, lane in enumerate(st.session_state.lanes) if lane['name'].upper() == "MK"]
    if mk_indices:
        st.subheader("MK 条带标记")
        mk_idx = st.selectbox(
            "选择要标记条带的MK胶道",
            options=mk_indices,
            format_func=lambda i: f"胶道 {i+1} (x={st.session_state.lanes[i]['x']})"
        )
        st.session_state.current_mk_idx = mk_idx
        selected_mk = st.selectbox("选择MK标准", mk_options, key="mk_std_select")
        if selected_mk == "自定义...":
            custom_name = st.text_input("自定义MK名称")
            custom_values = st.text_input("kDa值，逗号分隔", "250,150,100,75,50,37,25,20")
            if st.button("保存自定义"):
                if custom_name and custom_values:
                    vals = []
                    for v in custom_values.split(","):
                        v = v.strip()
                        if v:
                            if '.' in v:
                                vals.append(float(v))
                            else:
                                vals.append(int(v))
                    if vals:
                        add_mk_standard(custom_name, vals)
                        st.success(f"已保存 '{custom_name}'")
                        st.rerun()
            selected_mk = st.selectbox("选择MK标准 (更新后)", list(load_mk_standards().keys()))
        st.session_state.selected_mk_std = selected_mk

        current_bands = st.session_state.lanes[mk_idx]['bands']
        st.write(f"**当前已标记条带数: {len(current_bands)}**")
        if current_bands:
            kda_list = mk_standards.get(selected_mk, [])
            matched = kda_list[:len(current_bands)]
            st.write(f"对应kDa值: {matched}")
            for j, y in enumerate(current_bands):
                col1, col2 = st.columns([4, 1])
                col1.write(f"条带 {j+1}: y={y}")
                if col2.button("删除", key=f"del_band_{mk_idx}_{j}"):
                    st.session_state.lanes[mk_idx]['bands'].pop(j)
                    st.session_state.lanes[mk_idx]['bands'] = sorted(st.session_state.lanes[mk_idx]['bands'])
                    st.rerun()
        else:
            st.info("尚未标记条带")

        col1, col2 = st.columns(2)
        with col1:
            if st.button("开始标记条带"):
                st.session_state.mode = 'band'
                st.rerun()
        with col2:
            if st.button("清空所有条带"):
                st.session_state.lanes[mk_idx]['bands'] = []
                st.rerun()
    else:
        st.info("将某个胶道名称改为 'MK' 以启用条带标记")

    # ---- 标注样式参数 ----
    st.subheader("标注样式设置")
    col1, col2 = st.columns(2)
    with col1:
        label_offset = st.slider("胶道标注垂直偏移 (全局)", -200, 100, -10, 1)
        rotation = st.slider("胶道名称旋转角度 (°)", -90, 90, 0)
        font_choice = st.selectbox("字体", ["Arial", "Times New Roman", "Courier New", "SimHei"])
        lane_color = st.color_picker("胶道名称颜色", "#000000")
    with col2:
        band_offset = st.slider("kDa标注水平距离 (全局)", 0, 100, 15)
        font_size = st.slider("字体大小", 1, 48, 16, 1)
        band_color = st.color_picker("kDa标注颜色", "#000000")

    # ---- 预览控制 ----
    col_gen, col_auto = st.columns(2)
    with col_gen:
        if st.button("生成预览标注"):
            st.session_state.preview_auto = True
            try:
                if not st.session_state.lanes:
                    st.warning("请先标记胶道")
                else:
                    mk_std = st.session_state.selected_mk_std if mk_indices else ""
                    annotated = create_annotation_image(
                        pil_img.copy(),
                        st.session_state.lanes,
                        load_mk_standards(),
                        mk_std,
                        label_offset,
                        rotation,
                        font_choice,
                        band_offset,
                        font_size,
                        lane_color,
                        band_color,
                        offset_dict=st.session_state.offset_settings
                    )
                    st.session_state.preview_image = annotated
                    st.rerun()
            except Exception as e:
                st.error(f"预览失败: {e}")
    with col_auto:
        if st.session_state.preview_auto:
            st.success("🔄 自动预览已开启")
        else:
            st.info("点击「生成预览标注」以开启自动预览")

    # ---- 自动预览 ----
    if st.session_state.preview_auto and st.session_state.lanes:
        try:
            mk_std = st.session_state.selected_mk_std if mk_indices else ""
            annotated = create_annotation_image(
                pil_img.copy(),
                st.session_state.lanes,
                load_mk_standards(),
                mk_std,
                label_offset,
                rotation,
                font_choice,
                band_offset,
                font_size,
                lane_color,
                band_color,
                offset_dict=st.session_state.offset_settings
            )
            st.session_state.preview_image = annotated
        except Exception as e:
            st.error(f"自动预览更新失败: {e}")

    # ---- 精修与预览并排 ----
    col_fine, col_preview = st.columns([1, 1])
    with col_fine:
        with st.expander("🎯 精细调节标注位置（独立偏移）", expanded=True):
            st.write("分别调节胶道名称和每个MK条带的位置")
            # 胶道名称精修
            st.subheader("胶道名称偏移调节")
            lane_options = []
            for idx, lane in enumerate(st.session_state.lanes):
                label = f"胶道 {idx+1} 名称 ('{lane['name']}')"
                lane_options.append((idx, label))
            if not lane_options:
                st.info("暂无胶道")
            else:
                selected_idx = st.selectbox(
                    "选择胶道名称",
                    options=[i for i, _ in lane_options],
                    format_func=lambda i: dict(lane_options)[i],
                    key="fine_lane_select"
                )
                key = (selected_idx, 'label')
                current_dx, current_dy = st.session_state.offset_settings.get(key, (0, 0))
                dx = st.slider("X偏移 (像素, 正向右)", -100, 100, current_dx, 1, key=f"dx_lane_{selected_idx}")
                dy = st.slider("Y偏移 (像素, 正向下)", -100, 100, current_dy, 1, key=f"dy_lane_{selected_idx}")
                if (dx, dy) != (current_dx, current_dy):
                    st.session_state.offset_settings[key] = (dx, dy)
                    st.rerun()

            st.subheader("MK条带偏移调节")
            band_options = []
            for idx, lane in enumerate(st.session_state.lanes):
                if lane['name'].upper() == "MK" and lane.get('bands'):
                    kda_list = mk_standards.get(st.session_state.selected_mk_std, [])
                    sorted_bands = sorted(lane['bands'])
                    for bi, y in enumerate(sorted_bands):
                        if bi < len(kda_list):
                            label = f"MK胶道 {idx+1} - 条带 {bi+1} ({kda_list[bi]} kDa)"
                            band_options.append((idx, bi, label))
            if not band_options:
                st.info("暂无MK条带，请先标记MK条带")
            else:
                selected_band = st.selectbox(
                    "选择MK条带",
                    options=band_options,
                    format_func=lambda x: x[2],
                    key="fine_band_select"
                )
                idx, bi, _ = selected_band
                key = (idx, 'band', bi)
                current_dx, current_dy = st.session_state.offset_settings.get(key, (0, 0))
                dx = st.slider("X偏移 (像素, 正向右)", -100, 100, current_dx, 1, key=f"dx_band_{idx}_{bi}")
                dy = st.slider("Y偏移 (像素, 正向下)", -100, 100, current_dy, 1, key=f"dy_band_{idx}_{bi}")
                if (dx, dy) != (current_dx, current_dy):
                    st.session_state.offset_settings[key] = (dx, dy)
                    st.rerun()

    with col_preview:
        if st.session_state.preview_image is not None:
            img_w = st.session_state.preview_image.width
            st.image(st.session_state.preview_image, caption="标注预览", width=img_w)
            # 导出 SVG 和 PPTX
            try:
                mk_std = st.session_state.selected_mk_std if mk_indices else ""
                # SVG
                svg_str = generate_svg_overlay(
                    st.session_state.lanes,
                    load_mk_standards(),
                    mk_std,
                    label_offset,
                    rotation,
                    font_choice,
                    band_offset,
                    font_size,
                    lane_color,
                    band_color,
                    offset_dict=st.session_state.offset_settings,
                    img_width=pil_img.width,
                    img_height=pil_img.height,
                    background_img=pil_img
                )
                st.download_button(
                    label="📥 导出 SVG（含原图，标注可编辑）",
                    data=svg_str,
                    file_name="annotation_overlay.svg",
                    mime="image/svg+xml",
                    key="download_svg"
                )
                # PPTX
                if PPTX_AVAILABLE:
                    pptx_bytes = generate_pptx(
                        pil_img,
                        st.session_state.lanes,
                        load_mk_standards(),
                        mk_std,
                        label_offset,
                        rotation,
                        font_choice,
                        band_offset,
                        font_size,
                        lane_color,
                        band_color,
                        st.session_state.offset_settings
                    )
                    st.download_button(
                        label="📊 导出为 PPTX（含旋转，位置精确）",
                        data=pptx_bytes,
                        file_name="annotation.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        key="download_pptx"
                    )
                else:
                    st.info("💡 安装 python-pptx 以导出 PPTX：`pip install python-pptx`")
            except Exception as e:
                st.error(f"导出失败: {e}")
        else:
            st.info("请点击「生成预览标注」查看效果")

    # ---- 项目名称与保存 ----
    project_name = st.text_input("项目名称 (默认自动编号)", value=st.session_state.project_name)
    if not project_name:
        existing = list_projects()
        numbers = [int(p[1:]) for p in existing if p.startswith('p') and p[1:].isdigit()]
        next_num = max(numbers) + 1 if numbers else 1
        project_name = f"p{next_num}"
        st.info(f"自动使用名称: {project_name}")

    if st.button("保存标注项目"):
        try:
            if st.session_state.uploaded_img is None or not st.session_state.lanes:
                st.error("请先上传图片并标记胶道")
            else:
                mk_std = st.session_state.selected_mk_std if mk_indices else ""
                final_annotated = create_annotation_image(
                    pil_img.copy(),
                    st.session_state.lanes,
                    load_mk_standards(),
                    mk_std,
                    label_offset,
                    rotation,
                    font_choice,
                    band_offset,
                    font_size,
                    lane_color,
                    band_color,
                    offset_dict=st.session_state.offset_settings
                )
                data = {
                    "project_name": project_name,
                    "lanes": st.session_state.lanes,
                    "mk_type": mk_std,
                    "label_offset": label_offset,
                    "rotation": rotation,
                    "font": font_choice,
                    "band_offset": band_offset,
                    "font_size": font_size,
                    "lane_color": lane_color,
                    "band_color": band_color,
                    "offset_settings": st.session_state.offset_settings
                }
                save_project(project_name, data, st.session_state.uploaded_img, final_annotated)
                st.success(f"项目 '{project_name}' 已保存！")
                st.session_state.project_name = project_name
        except Exception as e:
            st.error(f"保存失败: {e}")

def manual_annotation():
    st.header("📤 上传胶图，点击标记胶道和条带")

    # ---- 状态初始化 ----
    if 'lanes' not in st.session_state:
        st.session_state.lanes = []
    if 'uploaded_img' not in st.session_state:
        st.session_state.uploaded_img = None
    if 'original_uploaded' not in st.session_state:
        st.session_state.original_uploaded = None
    if 'project_name' not in st.session_state:
        st.session_state.project_name = ""
    if 'mode' not in st.session_state:
        st.session_state.mode = 'lane'
    if 'current_mk_idx' not in st.session_state:
        st.session_state.current_mk_idx = -1
    if 'selected_mk_std' not in st.session_state:
        st.session_state.selected_mk_std = ""
    if 'preview_image' not in st.session_state:
        st.session_state.preview_image = None
    if 'preview_auto' not in st.session_state:
        st.session_state.preview_auto = False
    if 'offset_settings' not in st.session_state:
        st.session_state.offset_settings = {}
    if 'crop_left' not in st.session_state:
        st.session_state.crop_left = 0
        st.session_state.crop_top = 0
        st.session_state.crop_right = 0
        st.session_state.crop_bottom = 0

    # ---- 如果已有图片，直接进入标注界面 ----
    if st.session_state.uploaded_img is not None:
        show_annotation_interface(st.session_state.uploaded_img)
        return

    # ---- 否则显示输入方式 ----
    # 重置 MK 标准按钮
    col_reset, _ = st.columns([1, 5])
    with col_reset:
        if st.button("🔄 恢复默认 MK 标准"):
            reset_mk_standards()
            st.success("已恢复默认 MK 标准")
            st.rerun()

    # 文件上传
    st.markdown("### 📤 上传或拖拽图片")
    st.caption("💡 支持 JPG、PNG、TIF 格式，可直接拖拽截图或文件")
    uploaded_file = st.file_uploader(
        "点击选择或拖拽图片到此处",
        type=["jpg", "jpeg", "png", "tif", "tiff"],
        label_visibility="collapsed"
    )
    if uploaded_file is not None:
        pil_img = Image.open(uploaded_file)
        st.session_state.original_uploaded = pil_img
        st.session_state.uploaded_img = pil_img
        st.session_state.lanes = []
        st.session_state.preview_image = None
        st.success("✅ 图片加载成功！")
        st.rerun()

    # 粘贴图片
    st.markdown("---")
    col_paste, _ = st.columns([1, 3])
    with col_paste:
        try:
            from streamlit_paste_button import paste_image_button as pbutton
            paste_result = pbutton(
                label="📋 从剪贴板粘贴图片（复制后点击此处）",
                key="paste_button_main",
                errors="raise"
            )
            if paste_result.image_data is not None:
                pil_img = paste_result.image_data
                st.session_state.original_uploaded = pil_img
                st.session_state.uploaded_img = pil_img
                st.session_state.lanes = []
                st.session_state.preview_image = None
                st.success("✅ 图片粘贴成功！正在加载标注界面...")
                st.rerun()
        except ImportError:
            st.error("请先安装 streamlit-paste-button：`pip install streamlit-paste-button`")
        except Exception as e:
            st.error(f"粘贴失败: {e}")

def apply_crop(original_img, state):
    """根据 state 中的裁剪参数裁剪图像"""
    left = state.crop_left
    top = state.crop_top
    right = original_img.width - state.crop_right
    bottom = original_img.height - state.crop_bottom
    # 确保裁剪区域有效
    if right <= left:
        right = left + 1
    if bottom <= top:
        bottom = top + 1
    return original_img.crop((left, top, right, bottom))


def manage_projects():
    st.header("📂 管理已保存的胶图标注")
    projects = list_projects()
    if not projects:
        st.info("暂无保存的项目")
        return
    selected = st.selectbox("选择项目", projects)
    data = load_project(selected)
    if data is None:
        st.error("项目数据损坏或JSON格式错误，请删除该项目重新标注")
        if st.button("删除此项目"):
            if delete_project(selected):
                st.success("已删除")
                st.rerun()
            else:
                st.error("删除失败")
        return

    col1, col2 = st.columns(2)
    with col1:
        original = Image.open(data['original_image'])
        st.image(original, caption="原始图片", width='stretch')
    with col2:
        if os.path.exists(data.get('annotated_image', '')):
            annotated = Image.open(data['annotated_image'])
            st.image(annotated, caption="当前标注", width='stretch')
        else:
            st.warning("标注图片不存在")

    st.subheader("编辑参数")
    mk_standards = load_mk_standards()
    mk_options = list(mk_standards.keys())
    current_mk = data.get('mk_type', mk_options[0] if mk_options else "")

    col1, col2, col3 = st.columns(3)
    with col1:
        label_offset = st.number_input("垂直偏移", value=data.get('label_offset', -10), step=1)
        rotation = st.slider("旋转角度", -90, 90, data.get('rotation', 0))
        lane_color = st.color_picker("胶道颜色", data.get('lane_color', '#000000'))
    with col2:
        band_offset = st.number_input("kDa水平距离", value=data.get('band_offset', 15), step=1)
        font_size = st.number_input("字体大小", value=data.get('font_size', 16), step=1, min_value=1, max_value=48)
        band_color = st.color_picker("kDa颜色", data.get('band_color', '#000000'))
    with col3:
        font_choice = st.selectbox("字体", ["Arial", "Times New Roman", "Courier New", "SimHei"],
                                   index=["Arial","Times New Roman","Courier New","SimHei"].index(data.get('font','Arial')))
        mk_type = st.selectbox("MK类型", mk_options, index=mk_options.index(current_mk) if current_mk in mk_options else 0)

    st.subheader("编辑胶道和条带")
    lanes = data.get('lanes', [])
    new_lanes = []
    for idx, lane in enumerate(lanes):
        with st.expander(f"胶道 {idx+1} (x={lane['x']})", expanded=True):
            name = st.text_input("名称", value=lane.get('name',''), key=f"eman_{idx}")
            bands_str = ", ".join(str(b) for b in lane.get('bands', []))
            bands_input = st.text_input("条带y坐标 (逗号分隔)", value=bands_str, key=f"ebands_{idx}")
            bands = [int(b.strip()) for b in bands_input.split(",") if b.strip()]
            bands = sorted(bands)
            new_lanes.append({"x": lane['x'], "name": name, "bands": bands})

    offset_settings = data.get('offset_settings', {})

    if st.button("更新标注并保存"):
        original_img = Image.open(data['original_image'])
        updated = create_annotation_image(
            original_img.copy(), new_lanes, mk_standards, mk_type,
            label_offset, rotation, font_choice, band_offset, font_size,
            lane_color, band_color,
            offset_dict=offset_settings
        )
        data['lanes'] = new_lanes
        data['mk_type'] = mk_type
        data['label_offset'] = label_offset
        data['rotation'] = rotation
        data['font'] = font_choice
        data['band_offset'] = band_offset
        data['font_size'] = font_size
        data['lane_color'] = lane_color
        data['band_color'] = band_color
        data['offset_settings'] = offset_settings
        save_project(selected, data, original_img, updated)
        st.success("已更新")
        st.rerun()

    st.subheader("危险操作")
    if st.button("删除此项目"):
        if st.checkbox("确认删除？不可恢复！"):
            if delete_project(selected):
                st.success("已删除")
                st.rerun()
            else:
                st.error("删除失败")


if subpage == "手工标注":
    manual_annotation()
else:
    manage_projects()
